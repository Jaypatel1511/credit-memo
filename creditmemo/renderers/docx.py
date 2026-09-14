"""Render credit memo to Word .docx format."""
import datetime
import re

from creditmemo.data.schema import DealProfile
from creditmemo.renderers.markdown import render as render_markdown
from creditmemo.tables import block_to_rows, iter_segments
from creditmemo.text import strip_emphasis

#: The characters XML 1.0 forbids in character data, which is what a .docx
#: part is. Measured against python-docx rather than read off the spec: every
#: codepoint in this class was fed through ``Document.add_paragraph`` and
#: ``save`` and rejected, and ``\t``, ``\n``, ``\r`` and ``\x7f`` were fed
#: through and accepted. Control characters arrive in this package routinely —
#: a vertical tab or a form feed is what a paste out of a PDF leaves behind.
_XML_FORBIDDEN = re.compile(
    "[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


def _reject_control_characters(lines) -> None:
    """
    Raise a ValueError naming the character and the line that carries it.

    ``save_markdown`` accepts these characters and ``save_docx`` did not: the
    same DealProfile wrote one file and raised on the other with lxml's

        ValueError: All strings must be XML compatible: Unicode or ASCII,
        no NULL bytes or control characters

    which names no field, no line, no character and no file, and arrives after
    the caller has already been told the Markdown was fine. Nothing is
    sanitised — deleting or substituting a character is the silent alteration
    this release exists to remove — but the refusal now says what to fix.
    """
    for number, line in enumerate(lines, 1):
        match = _XML_FORBIDDEN.search(line)
        if match is None:
            continue
        char = match.group()
        raise ValueError(
            f"Word cannot store the character U+{ord(char):04X} that appears "
            f"at position {match.start()} of memo line {number}: {line!r}. "
            f"Control characters are routine in text pasted out of a PDF and "
            f"are not valid in a .docx; remove it from the input field it came "
            f"from. (save_markdown accepts it — only Word does not.)"
        )


#: Word style used for every generated table. "Table Grid" ships with the
#: default python-docx template, so no external template is required. The
#: README documents this style and the bold header row; both are gated.
TABLE_STYLE = "Table Grid"

#: Word's built-in bullet style. Unordered items are restyled into it because a
#: bullet glyph carries no information: whatever Word draws in the margin says
#: the same thing the "- " in the Markdown said.
#:
#: There is deliberately no numbered counterpart. See :func:`render`.
BULLET_STYLE = "List Bullet"


def _clean(text: str) -> str:
    """
    Strip Markdown emphasis markers from text bound for a Word run.

    Applied to every text branch. Before 0.2.1 only the plain-paragraph branch
    did this, so a bullet such as ``- **Total Assets:** $8.00MM`` reached Word
    as the literal characters ``**Total Assets:** $8.00MM`` — raw Markdown
    syntax, in the document handed to an Investment Committee.

    The rule itself lives in :mod:`creditmemo.text`, shared with the table-cell
    splitter. This branch used to delete every ``*`` while that one deleted
    only ``**``, so ``5 * 3`` survived in a Word table cell and arrived in a
    paragraph of the same document as ``5  3``.
    """
    return strip_emphasis(text)

#: The horizontal rule the Markdown renderer emits between sections. Matched
#: exactly, not by prefix: ``startswith("---")`` also matched underwriter prose
#: such as "--- see appendix" and replaced the whole sentence with a rule,
#: deleting the text from the Word document without a word on any stream.
RULE_LINE = "---"


def _add_rule(doc) -> None:
    """
    Append a horizontal rule as an empty paragraph with a bottom border.

    Before 0.2.1 this was a paragraph containing sixty literal underscores,
    which is a row of underscores in Word, not a rule: it does not span the
    column, does not follow the page margins, and is selectable text that lands
    in any copy-paste of the memo.
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    paragraph = doc.add_paragraph()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "auto")
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)


#: The spellings of ``DealProfile.prepared_date`` this package will read as a
#: date, tried in order.
#:
#: ``prepared_date`` is a free-form ``str`` and stays one — typing it is a
#: schema change and belongs to 0.3.0. Every format here is unambiguous. The
#: numeric day-first/month-first forms (``05/06/2026``) are deliberately
#: **not** here: reading one requires choosing between 5 June and 6 May, and
#: choosing is exactly the silent interpretation this package refuses
#: elsewhere. A value nothing here matches leaves the date properties unset —
#: it is not guessed at and it does not raise, because failing to stamp a
#: creation date is not a reason to refuse to write a memo.
PREPARED_DATE_FORMATS = (
    "%Y-%m-%d",
    "%Y/%m/%d",
    "%Y-%m-%dT%H:%M:%S",
    "%Y-%m-%d %H:%M:%S",
    "%B %d, %Y",
    "%b %d, %Y",
    "%B %d %Y",
    "%b %d %Y",
    "%d %B %Y",
    "%d %b %Y",
)


def parse_prepared_date(value) -> "datetime.datetime | None":
    """
    ``value`` as a datetime, or ``None`` if it is not a date this package reads.

    No timezone is invented. A value that carries an offset is normalised to
    UTC, which is a conversion and not a guess; a value that carries none stays
    naive. python-docx serialises either as ``...Z``, so a naive
    ``prepared_date`` is written to the file as though it were UTC — that is
    python-docx's encoding of a core property, not a claim this package makes,
    and it is the reason the *date* is the part worth stamping and the
    time-of-day is left at midnight.
    """
    if isinstance(value, datetime.datetime):
        parsed = value
    elif isinstance(value, datetime.date):
        parsed = datetime.datetime(value.year, value.month, value.day)
    elif isinstance(value, str):
        text = value.strip()
        parsed = None
        for fmt in PREPARED_DATE_FORMATS:
            try:
                parsed = datetime.datetime.strptime(text, fmt)
                break
            except ValueError:
                continue
    else:
        return None
    if parsed is None:
        return None
    if parsed.tzinfo is not None:
        parsed = parsed.astimezone(datetime.timezone.utc)
    return parsed


def _unset_core_datetime(properties, name: str) -> None:
    """
    Remove one ``dcterms`` datetime from the core properties part.

    python-docx has no way to *clear* a core datetime — ``cp.created = None``
    raises ``ValueError: property requires <type 'datetime.datetime'> object``,
    measured on python-docx 1.2.0 — and simply not setting it is not the same
    as leaving it unset: the default template ships a populated
    ``dcterms:created`` of ``2013-12-23T23:15Z``, so every memo that did not
    set it carried that date. Deleting the element leaves the property genuinely
    absent, which is what an unparseable ``prepared_date`` means.
    """
    from docx.oxml.ns import qn

    element = getattr(properties, "_element", None)
    if element is None:          # pragma: no cover - python-docx internals moved
        return
    node = element.find(qn("dcterms:%s" % name))
    if node is not None:
        element.remove(node)


def _set_core_properties(doc, deal: DealProfile) -> None:
    """
    Stamp the document's own metadata from the deal.

    Every ``.docx`` this package has ever written reported **Author
    python-docx** and **Created 2013-12-23**, because ``save_docx`` never
    touched core properties and python-docx's default template populates them.
    Measured on the published 0.2.1 wheel: ``author='python-docx'``,
    ``created=modified=2013-12-23T23:15Z``, ``title=''``.

    A credit memo goes into a loan file, and loan files get examined — by CDFI
    Fund compliance, by auditors, by the NMTC investor's counsel. Every document
    management system sorts and filters on that date, and it is false for every
    memo in the folder. It is also a fidelity defect under this package's own
    invariant: the metadata asserted a creation date that the ``prepared_date``
    printed on page one contradicted.

    ``author`` and ``title`` are set unconditionally: both sources are required
    ``str`` fields, and an empty one writes an empty property, which is honest —
    nobody claimed authorship — where ``python-docx`` is not. The dates are set
    only from a ``prepared_date`` this package can read; see
    :func:`parse_prepared_date`.

    ``modified`` is set to the same value as ``created`` rather than to the
    wall clock. The document is written once and never edited by this package,
    so its modification time *is* its creation time, and reaching for the clock
    would make two renders of the same deal differ in a way nothing in the deal
    accounts for.
    """
    properties = doc.core_properties
    properties.author = deal.prepared_by
    properties.title = deal.deal_name

    prepared = parse_prepared_date(deal.prepared_date)
    if prepared is None:
        _unset_core_datetime(properties, "created")
        _unset_core_datetime(properties, "modified")
        return
    properties.created = prepared
    properties.modified = prepared


def _add_table(doc, block) -> None:
    """
    Append one Markdown table block to ``doc`` as a real Word table, followed by
    an empty paragraph.

    That trailing paragraph is the one thing this renderer emits that is not a
    line of the Markdown: it is the blank line the Markdown has after every
    table, which the text path drops because a blank line adds no paragraph.
    Word runs a table and the next heading together without it.

    The ``not rows`` guard below is defensive and cannot fire through
    :func:`render`: ``iter_segments`` only calls a block a table when it has a
    delimiter rule, which needs at least two lines, and ``block_to_rows`` drops
    exactly one of them.
    """
    rows = block_to_rows(block)
    if not rows:
        return
    n_cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.style = TABLE_STYLE
    for r, cells in enumerate(rows):
        for c, text in enumerate(cells):
            cell = table.rows[r].cells[c]
            cell.text = text
            if r == 0:
                # Header row: bold every run python-docx created for the text.
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
    doc.add_paragraph()


def render(deal: DealProfile, path: str) -> None:
    """
    Render a complete credit memo as a Word .docx file.

    The document is a pure function of the Markdown the same deal produces:
    every heading, paragraph, list item, rule and table in the .docx comes from
    a line of that Markdown, and nothing carrying text is added that has no line
    behind it. The one addition is an empty spacer paragraph after each table —
    see :func:`_add_table`.

    ORDERED LISTS ARE NOT RESTYLED (R17). A line the Markdown begins with a
    number becomes a plain ``Normal`` paragraph whose text still carries that
    number. There is no ``List Number`` here, and no ``w:numPr``: nothing in
    this document asks Word to supply a number, because Word will not supply
    the number the Markdown states.

    An intermediate build of 0.2.1 did restyle them, and that is the defect
    this reverts — it never reached a released version. Word does not
    restart a numbered list on its own: every ``List Number`` paragraph
    python-docx creates resolves to one continuous numbering definition. A memo
    with a two-item ordered list in ``deal_summary`` and three ``conditions``
    reads 1,2 then 1,2,3 in the Markdown and printed the Conditions of Approval
    as **3, 4, 5** in Word. The same restyling deleted a caller's own literal
    number outright when a condition carried an embedded newline::

        conditions=["Receipt of appraisal",
                    "Payoff of the 2019 note\n3. Third-party report"]

    That build guarded against a nearby defect of 0.2.0 — a chronology written as
    ``2019. ...`` / ``2024. ...`` reaching the Investment Committee as items 1.
    and 2. — with a rule that only restyled a run numbered exactly 1..n. That
    rule is gone too, and with it the class: nothing is restyled, so nothing
    can be renumbered. The cost is that a genuine ordered list reads as ordinary
    paragraphs with their numbers intact, which is the direction this package
    fails in on purpose. Gated by G13.

    Unordered items keep ``List Bullet``: a bullet glyph carries no
    information, so letting Word draw it substitutes nothing for something.

    Before 0.2.1 the renderer built a cover block by hand — a ``Title``
    paragraph, an ``H1`` of the deal name and a four-row metadata table read
    straight off the ``DealProfile`` — and *then* parsed the Markdown, which
    carries all three of those things itself. The memo title, the deal name and
    the Fund/Prepared By/Date/IC Date block therefore appeared twice in every
    Word memo the package has ever produced, the second time in a different
    shape from the first. The hand-built block is gone; the first ``#`` heading
    simply takes Word's ``Title`` style.

    Requires python-docx: pip install python-docx

    Args:
        deal: DealProfile with all deal inputs
        path: Output path for the .docx file
    """
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        raise ImportError(
            "python-docx is required for Word output. "
            "Install it with: pip install python-docx"
        )

    doc = Document()
    _set_core_properties(doc, deal)

    md_content = render_markdown(deal)
    lines = md_content.split("\n")

    # Before anything is built, so the refusal names the problem instead of
    # arriving from lxml halfway through a document that is then not written.
    _reject_control_characters(lines)

    # iter_segments groups contiguous table lines into blocks and hands back
    # everything else line by line, in order. It is the same grouping the
    # conservation gates count with, so there is only one such rule in the
    # package.
    seen_title = False
    for kind, payload in iter_segments(lines):
        if kind == "table":
            _add_table(doc, payload)
            continue

        line = payload
        if line.startswith("# "):
            if not seen_title:
                # The memo's own title, in Word's Title style — one paragraph
                # for one Markdown heading, not a synthesised second copy.
                heading = doc.add_heading(_clean(line[2:]), 0)
                heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
                seen_title = True
            else:
                doc.add_heading(_clean(line[2:]), 1)
        elif line.startswith("## "):
            doc.add_heading(_clean(line[3:]), 2)
        elif line.startswith("### "):
            doc.add_heading(_clean(line[4:]), 3)
        elif line.strip() == RULE_LINE:
            _add_rule(doc)
        elif line.startswith("- ") or line.startswith("* "):
            doc.add_paragraph(_clean(line[2:]), style=BULLET_STYLE)
        elif line:
            clean = _clean(line)
            if clean.strip():
                doc.add_paragraph(clean)

    doc.save(path)
    print(f"Credit memo saved to {path}")
